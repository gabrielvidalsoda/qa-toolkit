# -*- coding: utf-8 -*-
"""
Gera a folha de execucao Excel do Plano de Testes Generico (Smoke / Sanidade).

O plano-fonte e CCM/QA/Plano-Testes-Generico-Smoke.md. Este script materializa
as checklists (seccao 5), o exemplo trabalhado (seccao 6) e a varredura de menu
(seccao 7) numa pasta de trabalho .xlsx com 3 abas:
  - Instrucoes : como preencher, legenda de status/severidade, contadores de progresso
  - Execucao   : 85 casos CT-GEN-*, com coluna "Resultado Real" para preencher
                 (listas suspensas em Status e Severidade)
  - Defeitos   : modelo de registo de bug (BUG-AAAAMMDD-nn)

Requisitos: Python 3 + openpyxl  (pip install openpyxl)
Uso:  python Scripts/gerar-plano-testes-generico-xlsx.py
Corre a partir da raiz do repositorio ou de qualquer pasta (resolve o caminho
de saida a partir da localizacao deste script).
"""
import datetime as dt
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(_RAIZ, "CCM", "QA", "Plano-Testes-Generico-Smoke-Execucao.xlsx")

wb = Workbook()

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
GRP_FILL = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

# =====================================================================
# Sheet 1 - Execucao
# =====================================================================
ws = wb.active
ws.title = "Execucao"

COLS = [
    ("ID", 20),
    ("Grupo", 12),
    ("Tela / Rota", 34),
    ("Acao / Passo", 46),
    ("Dados de teste", 22),
    ("Resultado Esperado", 60),
    ("Resultado Real", 46),
    ("Status", 15),
    ("Severidade", 11),
    ("ID Defeito", 16),
    ("Evidencia (print/link)", 24),
    ("Testador", 14),
    ("Data", 12),
    ("Notas", 30),
]
for i, (name, w) in enumerate(COLS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
    c.alignment = WRAP
    c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30


def R(id_, grp, tela, acao, dados, esperado):
    return [id_, grp, tela, acao, dados, esperado, "", "Nao executado", "", "", "", "", "", ""]


rows = []

nav = "Global - todas as telas"
rows += [
 R("CT-GEN-NAV-01", "NAV", nav, "Abrir a app na raiz / (Ctrl+F5)", "", "Layout carrega: menu a esquerda, header com \"Manta\", rodape institucional. Sem ecra branco. Console sem erro vermelho."),
 R("CT-GEN-NAV-02", "NAV", nav, "Expandir cada secao de topo do menu", "", "Sub-itens aparecem. Itens \"(em breve)\" ficam esmaecidos e nao navegam."),
 R("CT-GEN-NAV-03", "NAV", nav, "Clicar num item de menu com rota", "", "O conteudo muda, o URL muda, o breadcrumb reflete o caminho, o item fica marcado ativo."),
 R("CT-GEN-NAV-04", "NAV", nav, "Recarregar (F5) numa rota profunda", "", "A mesma tela recarrega - nao redireciona para / nem da 404."),
 R("CT-GEN-NAV-05", "NAV", nav, "Colar um URL de detalhe inexistente", "/medicamentos/factura/000000", "Mensagem amigavel de \"nao encontrado\". Sem stack trace, sem ecra branco, sem grade fantasma."),
 R("CT-GEN-NAV-06", "NAV", nav, "Botao \"Voltar\" do browser e o botao \"Voltar\" dentro da tela", "", "Ambos regressam a tela anterior preservando filtros e pagina da grade."),
 R("CT-GEN-NAV-07", "NAV", nav, "Colapsar / expandir o menu (toggle)", "", "Icones mantem-se; o estado colapsado nao parte o layout."),
 R("CT-GEN-NAV-08", "NAV", nav, "Cronometrar o carregamento da tela", "", "< ~5 s com VPN. Acima disso, anotar em Notas (nao e falha automatica)."),
]

pesq = "Tela de pesquisa (ex.: /medicamentos/factura)"
rows += [
 R("CT-GEN-PESQ-01", "PESQ", pesq, "Abrir a tela", "", "Painel de filtros visivel; grade em estado \"pre-pesquisa\". Sem chamada de dados automatica."),
 R("CT-GEN-PESQ-02", "PESQ", pesq, "Clicar e digitar em cada campo de texto", "", "O cursor entra; aceita digitacao; rotulo e placeholder em PT-PT."),
 R("CT-GEN-PESQ-03", "PESQ", pesq, "Digitar letras num campo numerico (NIF, codigo, no)", "abc", "Rejeita / aceita so digitos conforme a regra; mensagem de formato junto ao campo."),
 R("CT-GEN-PESQ-04", "PESQ", pesq, "Campo de data", "31/12/2024 digitado a mao", "Abre o date-picker ao clicar E aceita digitacao DD/MM/AAAA (defeito recorrente)."),
 R("CT-GEN-PESQ-05", "PESQ", pesq, "Abrir cada combobox (Estado, Tipo, Area...)", "", "Abre, mostra opcoes nao vazias, permite selecionar uma."),
 R("CT-GEN-PESQ-06", "PESQ", pesq, "Preencher 1 filtro valido e clicar \"Pesquisar\"", "", "Indicador de carregamento durante a chamada; a grade carrega resultados coerentes com o filtro."),
 R("CT-GEN-PESQ-07", "PESQ", pesq, "\"Pesquisar\" com todos os campos vazios", "", "Ou bloqueia com mensagem de criterio minimo, ou nao faz nada. Nunca rebenta."),
 R("CT-GEN-PESQ-08", "PESQ", pesq, "Valor invalido e \"Pesquisar\"", "FT@2026 / NIF 123 / data 99/99/9999", "Erro por campo; a consulta NAO e executada."),
 R("CT-GEN-PESQ-09", "PESQ", pesq, "Valor valido mas inexistente e \"Pesquisar\"", "", "Grade com \"sem resultados\". Nao e erro."),
 R("CT-GEN-PESQ-10", "PESQ", pesq, "Clicar \"Limpar\"", "", "Campos voltam ao estado inicial; erros somem; grade esvazia; URL perde os parametros de pesquisa."),
 R("CT-GEN-PESQ-11", "PESQ", pesq, "Premir Enter dentro de um campo", "", "Equivale a \"Pesquisar\" (ou nao faz nada). Nao submete um form vazio em silencio."),
 R("CT-GEN-PESQ-12", "PESQ", pesq, "Apos pesquisar, ver o URL e dar F5", "", "Os filtros estao na querystring e sobrevivem ao refresh."),
]

grid = "Grade de resultados (segue uma tela de pesquisa)"
rows += [
 R("CT-GEN-GRID-01", "GRID", grid, "Inspecionar os cabecalhos", "", "Legiveis, sem sobreposicao nem corte. Grupos \"Comunicado\"/\"Apurado\" presentes quando aplicavel."),
 R("CT-GEN-GRID-02", "GRID", grid, "Contar linhas por pagina", "", "10 por defeito. O total de registos / paginas e apresentado."),
 R("CT-GEN-GRID-03", "GRID", grid, "Paginacao", "", "Primeira / anterior / seguinte / ultima funcionam. Mudar o tamanho de pagina repoe a pagina 1."),
 R("CT-GEN-GRID-04", "GRID", grid, "Clicar num cabecalho ordenavel", "", "Alterna asc/desc, mostra a seta, repoe a pagina 1."),
 R("CT-GEN-GRID-05", "GRID", grid, "Grade com muitas colunas", "", "A grade ganha scroll horizontal proprio; a pagina nao."),
 R("CT-GEN-GRID-06", "GRID", grid, "Acao de linha", "", "Icone \"ver\" no fim da linha navega para o detalhe certo. Duplo-clique na linha NAO abre."),
 R("CT-GEN-GRID-07", "GRID", grid, "Linha \"TOTAIS\" (quando existe)", "", "Aparece no fim; os valores batem com a soma visivel."),
 R("CT-GEN-GRID-08", "GRID", grid, "Formato de valores", "", "Montantes 1.234,56; datas DD/MM/AAAA (defeito conhecido: 75.46 em vez de 75,46)."),
 R("CT-GEN-GRID-09", "GRID", grid, "Pesquisa sem resultados", "", "Texto de vazio explicito; nao uma grade so com cabecalhos e zero linhas."),
]

det = "Tela de detalhe (ex.: /medicamentos/factura/:id)"
rows += [
 R("CT-GEN-DET-01", "DET", det, "Abrir a partir da grade", "", "Vai para o registo certo; o URL contem o id."),
 R("CT-GEN-DET-02", "DET", det, "Procurar campos editaveis", "", "Nenhum. Sem inputs ativos nem botoes Guardar/Verificar/Conferir/Devolver/Anular."),
 R("CT-GEN-DET-03", "DET", det, "Listar as acoes da tela", "", "So navegacao (Voltar, Erros e diferencas, Ir para Fatura), visualizador de documento e icones \"ver\" das tabelas."),
 R("CT-GEN-DET-04", "DET", det, "Trocar de aba (Informacoes / Lotes / Prestacoes / Notas C-D...)", "", "O conteudo muda sem recarregar a pagina; a aba ativa fica marcada."),
 R("CT-GEN-DET-05", "DET", det, "Visualizador de documento", "", "Abre ao lado (PNG ou PDF). Sem documento -> botao desativado + mensagem de indisponivel."),
 R("CT-GEN-DET-06", "DET", det, "Icone \"ver\" numa sub-tabela", "", "Navega para o lote / prestacao correto (nao para uma aba interna)."),
 R("CT-GEN-DET-07", "DET", det, "Clicar \"Voltar\"", "", "Regressa a lista com filtros e pagina preservados."),
 R("CT-GEN-DET-08", "DET", det, "Colar o URL do detalhe diretamente", "", "Carrega na mesma, sem passar pela lista."),
]

form = "Formulario de criacao/edicao (ex.: Dados de Referencia) - ESCRITA"
rows += [
 R("CT-GEN-FORM-01", "FORM", form, "Clicar \"Novo\" / \"Criar\"", "", "Abre o formulario (modal ou tela); os campos obrigatorios estao assinalados."),
 R("CT-GEN-FORM-02", "FORM", form, "Guardar com obrigatorios vazios", "", "Erro por campo; NAO grava."),
 R("CT-GEN-FORM-03", "FORM", form, "Guardar com dados validos", "marcar QA-AAAAMMDD num campo livre", "Mensagem de sucesso; o registo aparece na lista."),
 R("CT-GEN-FORM-04", "FORM", form, "Testar tipo / limite em cada campo", "", "No de caracteres, digitos, formato de data validados, com a mensagem certa."),
 R("CT-GEN-FORM-05", "FORM", form, "Caracteres especiais / acentos / texto muito longo num campo de texto", "acn@#<>\"'; 500+ chars", "Aceite ou rejeitado com mensagem. Nunca 500."),
 R("CT-GEN-FORM-06", "FORM", form, "\"Cancelar\" / fechar o modal", "", "Nada e gravado; a lista fica igual."),
 R("CT-GEN-FORM-07", "FORM", form, "Editar o registo criado e dar F5", "", "A alteracao persiste."),
 R("CT-GEN-FORM-08", "FORM", form, "Apagar o registo criado", "", "Pede confirmacao; depois sai da lista."),
 R("CT-GEN-FORM-09", "FORM", form, "Clicar \"Guardar\" duas vezes rapido", "", "Nao cria dois registos."),
]

dash = "Dashboard / cockpit (ex.: /conferencia/cockpit)"
rows += [
 R("CT-GEN-DASH-01", "DASH", dash, "Abrir a tela", "", "Cartoes / graficos / contadores carregam, com indicador enquanto carregam."),
 R("CT-GEN-DASH-02", "DASH", dash, "Mexer nos filtros de topo (ciclo, area, periodo)", "", "Os numeros recalculam."),
 R("CT-GEN-DASH-03", "DASH", dash, "Clicar num cartao / linha clicavel", "", "Navega para o detalhe / lista correspondente."),
 R("CT-GEN-DASH-04", "DASH", dash, "Filtro sem dados", "", "Estado vazio explicito; nao grafico partido nem NaN."),
 R("CT-GEN-DASH-05", "DASH", dash, "Conferir coerencia", "", "Soma das partes = total apresentado."),
]

ui = "Transversal - todas as telas"
rows += [
 R("CT-GEN-UI-01", "UI", ui, "Varrer a tela (posicao/alinhamento/tamanho)", "", "Nada sobreposto, cortado, desalinhado ou fora do card."),
 R("CT-GEN-UI-02", "UI", ui, "Ler todos os textos", "", "Tudo em PT-PT: rotulos, botoes, mensagens, cabecalhos, tooltips. Sem chaves i18n cruas, sem ingles."),
 R("CT-GEN-UI-03", "UI", ui, "Comparar com outra tela do mesmo tipo", "", "Fontes, cores e espacamentos consistentes."),
 R("CT-GEN-UI-04", "UI", ui, "Botoes de acao", "", "O rotulo diz o que fazem; ficam desativados quando nao aplicavel (ex.: exportar sem pesquisa)."),
 R("CT-GEN-UI-05", "UI", ui, "Toda operacao assincrona", "", "Mostra spinner / barra de progresso e ele desaparece no fim."),
 R("CT-GEN-UI-06", "UI", ui, "Forcar erro de servico (5xx, timeout, VPN a cair)", "", "Snackbar de erro legivel no canto; tela nao fica branca nem com dados corrompidos; erro aparece na Console."),
 R("CT-GEN-UI-07", "UI", ui, "Atalhos e itens de menu", "", "Itens \"(em breve)\" nao navegam; \"logout\" no header leva a /login."),
 R("CT-GEN-UI-08", "UI", ui, "Redimensionar a janela para ~1280 e ~768 px", "", "O layout adapta; sem scroll horizontal na pagina (so dentro de grades/tabelas)."),
 R("CT-GEN-UI-09", "UI", ui, "Clicar em cada segmento do breadcrumb", "", "Navega para o nivel certo."),
]

xlsx = "Exportacao XLSX (grades de Consulta por Area)"
rows += [
 R("CT-GEN-XLSX-01", "XLSX", xlsx, "Ver o botao de exportar sem pesquisa ativa / com 0 resultados", "", "Desativado."),
 R("CT-GEN-XLSX-02", "XLSX", xlsx, "Exportar com resultados", "", "Descarrega um .xlsx (nao CSV nem PDF); nome do ficheiro coerente com a area."),
 R("CT-GEN-XLSX-03", "XLSX", xlsx, "Abrir o ficheiro", "", "Abre no Excel/LibreOffice; no de linhas ~= universo pesquisado; colunas fixas de exportacao."),
 R("CT-GEN-XLSX-04", "XLSX", xlsx, "Exportar com filtros aplicados", "", "O conteudo respeita os filtros ativos."),
]

ex = "/medicamentos/factura (dev:api, VPN ligada)"
rows += [
 R("CT-GEN-EXEMPLO-01", "EXEMPLO", ex, "Abrir http://localhost:5173/medicamentos/factura (Ctrl+F5).", "", "Tela \"Pesquisa Facturas\" carrega: filtros (No Factura, Data, No Processo, NIF Prestador, Codigo da Farmacia, Estado) + grade vazia. Console sem erro."),
 R("CT-GEN-EXEMPLO-02", "EXEMPLO", ex, "\"NIF Prestador\" = 501123083 -> \"Pesquisar\".", "NIF 501123083", "Spinner; a grade devolve ~44 faturas, todas com esse NIF; colunas \"Comunicado\" e \"Apurado\" visiveis."),
 R("CT-GEN-EXEMPLO-03", "EXEMPLO", ex, "Abrir \"Estado\", escolher \"Conferida Sem Erros\" -> \"Pesquisar\"; depois \"Conferida Com Erros\".", "Estado CSE / CCE", "URL ganha idEstado=14 (~17 registos). Trocar para \"Conferida Com Erros\" -> idEstado=15 (~27). Coluna Estado = texto do filtro."),
 R("CT-GEN-EXEMPLO-04", "EXEMPLO", ex, "Ir a pagina 2, a ultima, a primeira; mudar tamanho de pagina 10 -> 20.", "", "Navegacao correta entre paginas; mudar o tamanho repoe a pagina 1."),
 R("CT-GEN-EXEMPLO-05", "EXEMPLO", ex, "Clicar no cabecalho \"Data Factura\".", "", "Ordena por data (asc/desc); volta a pagina 1."),
 R("CT-GEN-EXEMPLO-06", "EXEMPLO", ex, "Clicar no icone \"ver\" no fim de uma linha.", "", "Navega para /medicamentos/factura/:id; abre o detalhe so-leitura com o visualizador do documento ao lado."),
 R("CT-GEN-EXEMPLO-07", "EXEMPLO", ex, "No detalhe, abrir a aba \"Lotes\"; clicar no icone \"ver\" de um lote.", "", "Navega para o detalhe do lote (nao uma aba interna)."),
 R("CT-GEN-EXEMPLO-08", "EXEMPLO", ex, "Clicar \"Voltar\".", "", "Regressa a grade de faturas com o filtro NIF + Estado e a pagina preservados."),
 R("CT-GEN-EXEMPLO-09", "EXEMPLO", ex, "Clicar \"Limpar\".", "", "Campos limpos, grade vazia, URL sem querystring."),
 R("CT-GEN-EXEMPLO-10", "EXEMPLO", ex, "(negativo) \"NIF Prestador\" = 123 -> \"Pesquisar\".", "NIF 123", "Erro junto ao campo (\"O NIF deve ter exatamente 9 digitos.\"); a consulta nao corre."),
]

rows += [
 R("CT-GEN-MENU-01", "MENU", "Inicio - /", "Abrir a secao e verificar o minimo", "", "Layout carrega, sem erro."),
 R("CT-GEN-MENU-02", "MENU", "Dados de Referencia - /dados-referencia/convencoes", "Abrir a secao e verificar o minimo", "ESCRITA", "Lista carrega; \"Novo\" abre formulario; a pesquisa filtra."),
 R("CT-GEN-MENU-03", "MENU", "Consulta por Area de Prestacao - /medicamentos/factura (repetir por area)", "Abrir a secao e verificar o minimo", "", "Pesquisa + grade + detalhe funcionam. 33 telas com a mesma forma (so leitura)."),
 R("CT-GEN-MENU-04", "MENU", "Gestao de Armazem - /armazem/registo-entrada", "Abrir a secao e verificar o minimo", "ESCRITA / massa partilhada", "Tela carrega; pesquisa; fluxo de escrita marcado QA-..."),
 R("CT-GEN-MENU-05", "MENU", "Controlo e Contabilizacao", "Abrir a secao e verificar o minimo", "", "So itens \"(em breve)\" - nada a testar."),
 R("CT-GEN-MENU-06", "MENU", "Conferencia - /conferencia/cockpit, /conferencia/worklist, /motor-regras/regras", "Abrir a secao e verificar o minimo", "ESCRITA", "Telas carregam; filtros; navegacao."),
 R("CT-GEN-MENU-07", "MENU", "Controlo de Qualidade - /controlo-qualidade/biblioteca-artefactos, /controlo-qualidade/dashboard-ciclo", "Abrir a secao e verificar o minimo", "ESCRITA", "Lista + dashboard carregam."),
 R("CT-GEN-MENU-08", "MENU", "Fecho", "Abrir a secao e verificar o minimo", "", "Item sem rota."),
 R("CT-GEN-MENU-09", "MENU", "Reclamacoes", "Abrir a secao e verificar o minimo", "", "Tudo \"(em breve)\"."),
 R("CT-GEN-MENU-10", "MENU", "Dashboards", "Abrir a secao e verificar o minimo", "", "Tudo \"(em breve)\"."),
 R("CT-GEN-MENU-11", "MENU", "Alarmistica", "Abrir a secao e verificar o minimo", "", "Dentro de Conferencia, \"(em breve)\"."),
]

prev_grp = None
band = False
for r_idx, row in enumerate(rows, start=2):
    if row[1] != prev_grp:
        band = not band
        prev_grp = row[1]
    for c_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.border = BORDER
        cell.alignment = WRAP if c_idx in (3, 4, 5, 6, 7, 14) else TOP
        if c_idx == 7:
            cell.fill = FILL_INPUT
        elif c_idx == 2 and band:
            cell.fill = GRP_FILL

last = len(rows) + 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:N{last}"

dv_status = DataValidation(type="list", formula1='"Nao executado,Passou,Falhou,Bloqueado,N/A"',
                           allow_blank=True, showErrorMessage=True, errorTitle="Valor invalido",
                           error="Use: Nao executado / Passou / Falhou / Bloqueado / N/A")
ws.add_data_validation(dv_status)
dv_status.add(f"H2:H{last}")

dv_sev = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Severidade",
                        error="1 = tudo / 2 = funcional + quebra de layout / 3 = funcional severo")
ws.add_data_validation(dv_sev)
dv_sev.add(f"I2:I{last}")

# =====================================================================
# Sheet 2 - Defeitos
# =====================================================================
wd = wb.create_sheet("Defeitos")
DCOLS = [
 ("ID Defeito", 18), ("Titulo", 40), ("Tela / Rota", 30), ("Passos para reproduzir", 46),
 ("Resultado Esperado", 36), ("Resultado Observado", 36), ("Evidencia (print/link)", 24),
 ("Severidade", 11), ("Ambiente", 22), ("Testador", 14), ("Data", 12),
 ("Estado", 16), ("Notas", 28),
]
for i, (name, w) in enumerate(DCOLS, start=1):
    c = wd.cell(row=1, column=i, value=name)
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
    c.alignment = WRAP
    c.border = BORDER
    wd.column_dimensions[get_column_letter(i)].width = w
wd.row_dimensions[1].height = 30
for r_idx in range(2, 42):
    for c_idx in range(1, len(DCOLS) + 1):
        cell = wd.cell(row=r_idx, column=c_idx, value="")
        cell.border = BORDER
        cell.alignment = WRAP
        cell.fill = FILL_INPUT
    wd.cell(row=r_idx, column=1, value=f"BUG-{dt.date.today():%Y%m%d}-{r_idx-1:02d}")
wd.freeze_panes = "A2"
wd.auto_filter.ref = "A1:M41"
dv_sev2 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
wd.add_data_validation(dv_sev2)
dv_sev2.add("H2:H41")
dv_estado = DataValidation(type="list",
                           formula1='"Aberto,Em analise,Corrigido,Fechado,Duplicado,Nao e defeito"',
                           allow_blank=True)
wd.add_data_validation(dv_estado)
dv_estado.add("L2:L41")

# =====================================================================
# Sheet 3 - Instrucoes (first)
# =====================================================================
wi = wb.create_sheet("Instrucoes", 0)
wi.column_dimensions["A"].width = 4
wi.column_dimensions["B"].width = 30
wi.column_dimensions["C"].width = 95
H = Font(bold=True, size=13, color="1F3864")
SUB = Font(bold=True, size=11)


def put(row, b="", c="", bold_b=False, header=False):
    cb = wi.cell(row=row, column=2, value=b)
    cc = wi.cell(row=row, column=3, value=c)
    cc.alignment = WRAP
    if header:
        cb.font = H
    elif bold_b:
        cb.font = SUB
    return row + 1


r = 2
r = put(r, "Plano de Testes Generico - Folha de Execucao", header=True)
r += 1
r = put(r, "Documento-fonte", "CCM/QA/Plano-Testes-Generico-Smoke.md - ler antes de comecar.")
r = put(r, "Objetivo", "Smoke / sanidade do ccmsns-ui: a tela abre, os botoes respondem, a pesquisa devolve resultado coerente, a navegacao leva onde diz.")
r += 1
r = put(r, "Como preencher", bold_b=True)
r = put(r, "1", "Na aba \"Execucao\", para cada linha: executar a acao, escrever o que aconteceu na coluna amarela \"Resultado Real\".")
r = put(r, "2", "Definir \"Status\": Passou / Falhou / Bloqueado / N/A (lista suspensa).")
r = put(r, "3", "Se \"Falhou\": preencher \"Severidade\" (1/2/3), tirar print, e abrir uma linha na aba \"Defeitos\" com o ID em \"ID Defeito\".")
r = put(r, "4", "\"Bloqueado\" = impossivel testar (sem massa, tela nao abre, sem 2o utilizador). Nao usar para \"nao sei que valor por\".")
r = put(r, "5", "Preencher Testador e Data. Usar \"Notas\" para massa usada, intermitencias, ambiente.")
r = put(r, "6", "Falha sem print degrada para observacao - regista-se, mas nao conta como falha confirmada.")
r += 1
r = put(r, "Status possiveis", bold_b=True)
r = put(r, "Nao executado", "Ainda nao testado (valor inicial).")
r = put(r, "Passou", "Resultado real = resultado esperado.")
r = put(r, "Falhou", "Desvio do esperado. Exige print + Severidade + linha em \"Defeitos\".")
r = put(r, "Bloqueado", "Nao foi possivel testar. Razao obrigatoria em Notas.")
r = put(r, "N/A", "O caso nao se aplica a esta tela (ex.: tela sem exportacao).")
r += 1
r = put(r, "Escala de severidade", bold_b=True)
r = put(r, "1", "Reporta qualquer desvio: funcionalidade, layout, estilo/CSS, acessibilidade basica, e sugestoes de melhoria.")
r = put(r, "2  (padrao)", "Reporta funcionalidade e quebra de layout - sobreposto, cortado, ilegivel. Ignora cosmetico e melhorias.")
r = put(r, "3", "So funcionalidade severa: nao faz o que promete, grava o que nao devia, rebenta com 5xx, perde dados, tranca o fluxo sem saida.")
r += 1
r = put(r, "Ambiente", bold_b=True)
r = put(r, "URL", "npm run dev:api -> http://localhost:5173  (mock: dev:mock, mesma URL; preview: :4174)")
r = put(r, "Pre-requisitos", "VPN SPMS ligada para dados reais. Sem login (header mostra \"Manta\"). Ctrl+F5 antes de cada tela. Chrome + DevTools (Console/Network).")
r = put(r, "Grade vazia?", "Pode ser filtro sem dados em DEV. Trocar o filtro antes de reportar. Massa conhecida: seccao 9 do documento-fonte.")
r += 1
r = put(r, "Progresso", bold_b=True)
prog = [
 ("Total de casos", "=COUNTA(Execucao!A2:A200)"),
 ("Executados", '=COUNTIF(Execucao!H2:H200,"Passou")+COUNTIF(Execucao!H2:H200,"Falhou")+COUNTIF(Execucao!H2:H200,"Bloqueado")+COUNTIF(Execucao!H2:H200,"N/A")'),
 ("Passou", '=COUNTIF(Execucao!H2:H200,"Passou")'),
 ("Falhou", '=COUNTIF(Execucao!H2:H200,"Falhou")'),
 ("Bloqueado", '=COUNTIF(Execucao!H2:H200,"Bloqueado")'),
 ("N/A", '=COUNTIF(Execucao!H2:H200,"N/A")'),
 ("Falhas sev. 3", '=COUNTIFS(Execucao!H2:H200,"Falhou",Execucao!I2:I200,3)'),
]
for label, formula in prog:
    wi.cell(row=r, column=2, value=label)
    wi.cell(row=r, column=3, value=formula)
    r += 1

wb.save(OUT)
print("OK ->", OUT, "| casos:", len(rows))
